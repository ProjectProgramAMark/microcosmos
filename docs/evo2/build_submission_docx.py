"""Render the canonical Evo² Markdown report as an image-embedded DOCX."""

from __future__ import annotations

import argparse
import base64
import mimetypes
from pathlib import Path

import markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt
from html2docx import html2docx


def _inline_images(html: str, source_dir: Path) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for image in soup.find_all("img"):
        source = image.get("src")
        if not source or source.startswith(("data:", "http://", "https://")):
            continue
        path = (source_dir / source).resolve()
        mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        image["src"] = f"data:{mime};base64,{encoded}"
    return str(soup)


def _style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)

    for name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True

    for table in document.tables:
        table.style = "Table Grid"
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)

    document.core_properties.title = (
        "Evo²-Ecosystem: Recursive Discovery of Heredity Policies"
    )
    document.core_properties.author = "Mark Moussa"
    document.core_properties.subject = "Sakana AI pre-interview project"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()

    source = arguments.source.resolve()
    html = markdown.markdown(
        source.read_text(encoding="utf-8"),
        extensions=("tables", "fenced_code", "sane_lists"),
    )
    html = _inline_images(html, source.parent)
    buffer = html2docx(html, title="Evo²-Ecosystem")
    document = Document(buffer)
    _style(document)
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(arguments.output)


if __name__ == "__main__":
    main()
